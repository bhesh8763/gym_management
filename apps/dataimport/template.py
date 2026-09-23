"""
Builds the downloadable .xlsx import template: a Read Me tab with the rules
plus one tab per data type, headers styled and frozen, one example row each.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .schema import SHEETS

HEADER_FILL = PatternFill(fill_type='solid', start_color='1F4E79')

_README = [
    ('FitCore — bulk data import template', 'title'),
    ('', 'text'),
    ('How to use', 'heading'),
    ('1. Fill in as many sheets as you need — sheets you leave out are simply skipped.', 'text'),
    ('2. Keep row 1 (the headers) exactly as it is. Column order does not matter, and extra columns are ignored.', 'text'),
    ('3. Upload the file on the Import Data page, review the preview, then confirm the import.', 'text'),
    ('', 'text'),
    ('Rules', 'heading'),
    ('• Dates: YYYY-MM-DD (01/03/2026 also works). Times: HH:MM in 24h form, e.g. 07:30 or 19:00.', 'text'),
    ('• Yes/No columns accept Yes, No, true, false, 1 and 0.', 'text'),
    ('• Multiple values (features, specializations, plans): separate them with a semicolon, e.g. "A; B".', 'text'),
    ('• Email is the unique key for people, receipts for payments, locker numbers for lockers.', 'text'),
    ('• Records that already exist are skipped, never duplicated — re-uploading the same file is safe.', 'text'),
    ('• Sheets import in a fixed order: plans → people → offers → memberships → payments → attendance → equipment → lockers → assignments → progress.', 'text'),
    ('• Leave Password blank to create an account that cannot sign in until its password is reset.', 'text'),
    ('• Rows with errors block the whole import — fix them and re-upload. Nothing is saved until you confirm.', 'text'),
    ('• Limits: .xlsx only, up to 4.5 MB, up to 20,000 rows per sheet.', 'text'),
    ('', 'text'),
    ('Sheets', 'heading'),
]


def build_template_bytes():
    """Return the template workbook as .xlsx bytes."""
    workbook = Workbook()
    readme = workbook.active
    readme.title = 'Read Me'
    for row_number, (text, style) in enumerate(_README, start=1):
        cell = readme.cell(row=row_number, column=1, value=text)
        if style == 'title':
            cell.font = Font(bold=True, size=16, color='1F4E79')
        elif style == 'heading':
            cell.font = Font(bold=True, size=12, color='1F4E79')
    sheet_line = len(_README) + 1
    for sheet_cls in SHEETS:
        cell = readme.cell(
            row=sheet_line, column=1,
            value=f'• {sheet_cls.name} — {sheet_cls.description}',
        )
        cell.font = Font(size=11)
        sheet_line += 1
    readme.column_dimensions['A'].width = 130

    for sheet_cls in SHEETS:
        worksheet = workbook.create_sheet(sheet_cls.name)
        for position, column in enumerate(sheet_cls.columns, start=1):
            header = worksheet.cell(row=1, column=position, value=column.label)
            header.font = Font(bold=True, color='FFFFFF')
            header.fill = HEADER_FILL
            header.alignment = Alignment(horizontal='center', vertical='center')
            example = str(sheet_cls.example.get(column.key, ''))
            width = max(len(column.label) + 4, len(example) + 4, 12)
            worksheet.column_dimensions[header.column_letter].width = min(width, 42)
            worksheet.cell(row=2, column=position, value=example)
        worksheet.freeze_panes = 'A2'

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
