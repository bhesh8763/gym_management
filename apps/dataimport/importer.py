"""
Upload → validate → preview → commit pipeline for the bulk workbook import.

Two passes:

  Pass A (read-only) — parse every sheet, validate rows, catch in-file
      duplicates and rows the database already has. Anything this file will
      create is registered on the context as "planned", so cross-references
      inside one workbook validate before a single row is written.

  Pass B (commit) — only when dry_run is off and Pass A found no errors:
      create everything inside one transaction. Any failure rolls the whole
      import back, so the database either receives the complete file or
      nothing at all.
"""
import logging

from django.db import transaction
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.lockers.models import Locker
from apps.memberships.models import Membership, MembershipPlan, Offer

from .schema import SHEETS, RowError, as_text, norm_header

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 4_500_000          # the project caps request bodies at 5 MB
MAX_ROWS_PER_SHEET = 20_000
MAX_ERRORS_PER_SHEET = 50
MAX_SKIPS_SHOWN = 15
ALLOWED_EXTENSIONS = ('.xlsx', '.xlsm')


class _Abort(Exception):
    """Raised inside the commit transaction to force a full rollback."""


# ─────────────────────────────────────────────────────────────────────────────
# Shared import context
# ─────────────────────────────────────────────────────────────────────────────

class Ctx:
    """Everything one import run needs to share between its sheets."""

    def __init__(self, actor=None, default_password=None):
        self.actor = actor                    # the user running the import
        self.default_password = default_password
        self.passwordless_accounts = 0

        self.users = {}                       # email -> User (in the database)
        self.planned = {}                     # email -> role (this file creates it)
        self.plans = {}                       # plan name -> MembershipPlan
        self.planned_plans = set()            # plan names this file creates
        self.lockers = {}                     # locker number -> Locker
        self.planned_lockers = set()
        self.offers = {}                      # offer name -> Offer
        self.planned_offers = set()
        self.memberships = {}                 # (email, plan) -> Membership (written)
        self.planned_memberships = set()      # (email, plan) this file creates

        self.seen = {}                        # (sheet key, identity) -> first row

        # In-file "only one active …" bookkeeping, checked before the database.
        self.active_lockers = set()
        self.active_member_lockers = set()
        self.active_trainee_members = set()

    def load(self):
        """Cache every lookup table used by validation (four queries total)."""
        self.users = {user.email.lower(): user for user in User.objects.all()}
        self.plans = {plan.name.lower(): plan for plan in MembershipPlan.objects.all()}
        self.lockers = {locker.locker_number.lower(): locker for locker in Locker.objects.all()}
        self.offers = {offer.name.lower(): offer for offer in Offer.objects.all()}

    # ── Lookups used while validating ──────────────────────────────────────

    def resolve_user(self, email, column, role=None, required=True):
        """Find an account, honouring accounts this same file will create.

        Returns the User, or None when the account only exists on paper
        (created by this file — it becomes a real User during the commit).
        """
        key = (email or '').strip().lower()
        if not key:
            if required:
                raise RowError(column, 'This field is required.')
            return None
        user = self.users.get(key)
        if user is None:
            planned_role = self.planned.get(key)
            if planned_role is None:
                raise RowError(
                    column,
                    f'No account found for "{key}" — add it to the Members, Staff or '
                    f'Trainers sheet (or create it in the app) first.',
                )
            if role and planned_role != role:
                raise RowError(
                    column,
                    f'"{key}" is being imported as a {planned_role} account, but this '
                    f'row needs a {role} account.',
                )
            return None
        if role and user.role != role:
            raise RowError(column, f'"{key}" is a {user.role} account, not a {role} account.')
        return user

    def require_user(self, email, column, role=None):
        """Commit-time variant: the account must be a real User by now."""
        user = self.resolve_user(email, column, role=role)
        if user is None:
            raise RowError(
                column,
                f'"{email}" should have been created by this import but was not — '
                f'check the Members, Staff or Trainers sheet.',
            )
        return user

    def resolve_plan(self, name, column):
        key = (name or '').strip().lower()
        if not key:
            raise RowError(column, 'This field is required.')
        plan = self.plans.get(key)
        if plan is None:
            if key not in self.planned_plans:
                raise RowError(
                    column,
                    f'No membership plan named "{name}" — add it to the Membership '
                    f'Plans sheet first.',
                )
            return None
        return plan

    def require_plan(self, name, column):
        plan = self.resolve_plan(name, column)
        if plan is None:
            raise RowError(column, f'The plan "{name}" should have been created by this import but was not.')
        return plan

    def resolve_locker(self, number, column):
        key = (number or '').strip().lower()
        if not key:
            raise RowError(column, 'This field is required.')
        locker = self.lockers.get(key)
        if locker is None and key not in self.planned_lockers:
            raise RowError(
                column,
                f'No locker numbered "{number}" — add it to the Lockers sheet first.',
            )
        return locker

    def require_locker(self, number, column):
        locker = self.resolve_locker(number, column)
        if locker is None:
            raise RowError(column, f'Locker "{number}" should have been created by this import but was not.')
        return locker

    def resolve_offer(self, name, column):
        key = (name or '').strip().lower()
        if not key:
            raise RowError(column, 'This field is required.')
        offer = self.offers.get(key)
        if offer is None:
            if key not in self.planned_offers:
                raise RowError(
                    column,
                    f'No offer named "{name}" — add it to the Offers sheet first.',
                )
            return None
        return offer

    def require_offer(self, name, column):
        offer = self.resolve_offer(name, column)
        if offer is None:
            raise RowError(column, f'The offer "{name}" should have been created by this import but was not.')
        return offer

    def membership_ref_valid(self, email, plan_name):
        """Does this member have (or get) a membership on that plan?"""
        key = (email.strip().lower(), plan_name.strip().lower())
        if key in self.memberships or key in self.planned_memberships:
            return True
        return Membership.objects.filter(
            member__email__iexact=key[0], plan__name__iexact=key[1],
        ).exists()

    def require_membership(self, email, plan_name, column):
        key = (email.strip().lower(), plan_name.strip().lower())
        membership = self.memberships.get(key)
        if membership is None:
            membership = Membership.objects.filter(
                member__email__iexact=key[0], plan__name__iexact=key[1],
            ).order_by('-start_date').first()
            if membership is None:
                raise RowError(column, f'No "{plan_name}" membership found for {key[0]}.')
            self.memberships[key] = membership
        return membership


# ─────────────────────────────────────────────────────────────────────────────
# Report helpers
# ─────────────────────────────────────────────────────────────────────────────

def _new_entry(sheet_cls):
    return {
        'key': sheet_cls.key,
        'name': sheet_cls.name,
        'present': False,
        'rows': 0,
        'to_create': 0,
        'skipped': 0,
        'error_count': 0,
        'errors': [],
        'errors_truncated': False,
        'skips': [],
        'ignored_columns': [],
    }


def _add_error(entry, row_number, error):
    """Record a row error (list capped, counter uncapped)."""
    entry['error_count'] += 1
    if len(entry['errors']) >= MAX_ERRORS_PER_SHEET:
        entry['errors_truncated'] = True
        return
    entry['errors'].append({
        'row': row_number,
        'column': error.column,
        'message': error.message,
    })


def _add_skip(entry, row_number, reason):
    if len(entry['skips']) < MAX_SKIPS_SHOWN:
        entry['skips'].append({'row': row_number, 'reason': reason})


def _summarize(report):
    totals = {'rows': 0, 'to_create': 0, 'skipped': 0, 'errors': 0}
    for entry in report['sheets']:
        totals['rows'] += entry['rows']
        totals['to_create'] += entry['to_create']
        totals['skipped'] += entry['skipped']
        totals['errors'] += entry['error_count']
    report['totals'] = totals


def _find_sheet_title(sheet_cls, title_index):
    for candidate in (sheet_cls.name, *sheet_cls.aliases):
        title = title_index.get(norm_header(candidate))
        if title:
            return title
    return None


# ─────────────────────────────────────────────────────────────────────────────
# Pass A — validation (no writes)
# ─────────────────────────────────────────────────────────────────────────────

def _validate_sheet(sheet_cls, workbook, title_index, ctx):
    """Validate one sheet; returns (report entry, rows that would be created)."""
    entry = _new_entry(sheet_cls)

    title = _find_sheet_title(sheet_cls, title_index)
    if title is None:
        return entry, []
    entry['present'] = True
    worksheet = workbook[title]

    rows_iter = worksheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None:
        _add_error(
            entry, 1,
            RowError('', 'This sheet is empty — row 1 must contain the column headers.'),
        )
        return entry, []

    # Map header cells to column keys. Unknown columns are reported, not fatal.
    header_map = sheet_cls.header_map()
    column_index = {}
    seen_keys = set()
    for position, cell in enumerate(header_row):
        key = header_map.get(norm_header(cell))
        if key is None:
            if as_text(cell):
                entry['ignored_columns'].append(as_text(cell))
            continue
        if key in seen_keys:
            continue
        seen_keys.add(key)
        column_index[position] = key

    missing = [
        column.label for column in sheet_cls.columns
        if column.required and column.key not in seen_keys
    ]
    if missing:
        _add_error(
            entry, 1,
            RowError(
                ', '.join(missing),
                'Missing required column(s) — download the template and keep its header row.',
            ),
        )
        return entry, []

    # Data rows (row 1 is the header, so openpyxl's 2nd row onwards).
    to_create = []
    row_number = 1
    for values in rows_iter:
        row_number += 1
        if all(value is None or as_text(value) == '' for value in values):
            continue
        entry['rows'] += 1
        if len(to_create) >= MAX_ROWS_PER_SHEET:
            _add_error(
                entry, row_number,
                RowError('', f'Too many rows — a sheet is limited to {MAX_ROWS_PER_SHEET:,}.'),
            )
            break

        raw = {
            key: values[position]
            for position, key in column_index.items()
            if position < len(values)
        }
        try:
            data = sheet_cls.clean(row_number, raw, ctx)

            identity = sheet_cls.identity(data)
            if identity is not None:
                marker = (sheet_cls.key, identity)
                first_row = ctx.seen.get(marker)
                if first_row is not None:
                    raise RowError(
                        sheet_cls.columns[0].label,
                        f'Duplicate of row {first_row} in this sheet.',
                    )
                ctx.seen[marker] = row_number

            reason = sheet_cls.existing(data, ctx)
        except RowError as error:
            _add_error(entry, row_number, error)
            continue

        if reason:
            entry['skipped'] += 1
            _add_skip(entry, row_number, reason)
        else:
            to_create.append((row_number, data))

    entry['to_create'] = len(to_create)
    ctx.passwordless_accounts += sheet_cls.default_password_count(to_create, ctx)
    return entry, to_create


# ─────────────────────────────────────────────────────────────────────────────
# Pass B — commit (only reached with zero validation errors)
# ─────────────────────────────────────────────────────────────────────────────

def _commit_sheet(sheet_cls, entry, rows, ctx):
    """Create the rows for one sheet. Returns True when something failed."""
    failed = False
    for row_number, data in rows:
        try:
            sheet_cls.insert(data, ctx)
        except RowError as error:
            _add_error(entry, row_number, error)
            failed = True
    return failed


# ─────────────────────────────────────────────────────────────────────────────
# Entry point
# ─────────────────────────────────────────────────────────────────────────────

def run_import(upload, *, actor, dry_run=True, default_password=None):
    """Validate (dry_run=True) or import an uploaded workbook.

    Always returns a report dict; `report['ok']` is True only when the file
    parsed cleanly and every row is importable.
    """
    report = {
        'ok': False,
        'dry_run': dry_run,
        'file': getattr(upload, 'name', None) or 'workbook.xlsx',
        'error': None,
        'sheets': [],
        'accounts_using_default_password': 0,
        'accounts_without_password': 0,
        'totals': {'rows': 0, 'to_create': 0, 'skipped': 0, 'errors': 0},
    }
    ctx = Ctx(actor=actor, default_password=default_password)

    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        logger.exception('Could not parse uploaded workbook %s', report['file'])
        report['error'] = (
            'This file could not be read as an Excel workbook. '
            'Save it as .xlsx (not .csv, .xls or a password-protected file) and try again.'
        )
        return report

    try:
        ctx.load()
        # normalized tab name -> actual tab title
        title_index = {norm_header(ws.title): ws.title for ws in workbook.worksheets}

        plan = []  # (entry, sheet_cls, to_create) in import order
        found_any = False
        for sheet_cls in SHEETS:
            entry, to_create = _validate_sheet(sheet_cls, workbook, title_index, ctx)
            report['sheets'].append(entry)
            found_any = found_any or entry['present']
            plan.append((entry, sheet_cls, to_create))

        if not found_any:
            report['error'] = (
                'No supported sheets were found in this workbook. Expected tab names '
                'such as: ' + ', '.join(sheet.name for sheet in SHEETS) + '.'
            )
        elif not dry_run:
            if any(entry['error_count'] for entry, _, _ in plan):
                # Validation failed — Pass B never runs, so nothing is written.
                report['error'] = 'Nothing was imported — fix the errors below and try again.'
            else:
                # One transaction for the whole file — all or nothing.
                try:
                    with transaction.atomic():
                        failed = False
                        for entry, sheet_cls, to_create in plan:
                            if _commit_sheet(sheet_cls, entry, to_create, ctx):
                                failed = True
                                break
                        if failed:
                            raise _Abort
                except _Abort:
                    report['error'] = (
                        'Import failed — no changes were saved. Fix the errors below '
                        'and try again.'
                    )

    except Exception:
        logger.exception('Bulk import failed unexpectedly')
        report['error'] = 'The import failed unexpectedly — no changes were saved.'
    finally:
        workbook.close()

    if default_password:
        report['accounts_using_default_password'] = ctx.passwordless_accounts
        report['accounts_without_password'] = 0
    else:
        report['accounts_using_default_password'] = 0
        report['accounts_without_password'] = ctx.passwordless_accounts
    _summarize(report)
    report['ok'] = report['error'] is None and report['totals']['errors'] == 0
    return report
